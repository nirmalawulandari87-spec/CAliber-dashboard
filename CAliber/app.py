from flask import Flask, request, jsonify
from flask_cors import CORS
import pandas as pd
import numpy as np
import pickle
import glob
import os
import shap
from sklearn.feature_extraction.text import TfidfVectorizer
from sklearn.metrics.pairwise import cosine_similarity
from pptx import Presentation

app = Flask(__name__)
CORS(app) # Mengizinkan frontend untuk akses API ini

# ----------------------------------------------------
# 1. LOAD AI 1 (ISOLATION FOREST MODELS)
# ----------------------------------------------------
models = {}
cols = {}
MODEL_DIR = "models"
for file in glob.glob(f"{MODEL_DIR}/ai1_model_*.pkl"):
    machine = os.path.basename(file).replace("ai1_model_", "").replace(".pkl", "")
    with open(file, 'rb') as f:
        models[machine] = pickle.load(f)
    with open(f"{MODEL_DIR}/cols_{machine}.pkl", 'rb') as f:
        cols[machine] = pickle.load(f)

# ----------------------------------------------------
# 2. BUILD AI 3 KNOWLEDGE BASE DARI PPT
# ----------------------------------------------------
knowledge_base = []
RCA_DIR = "data_rca"
print("Membangun Knowledge Base AI 3 dari file RCA (PPT)...")
for file in glob.glob(f"{RCA_DIR}/*.pptx"):
    # Deteksi mesin dari nama file
    machine_name = "UNKNOWN"
    if "PU-2101B" in file: machine_name = "PU-2101B"
    elif "KO-3201" in file: machine_name = "KO-3201"
    elif "PM-4405B" in file: machine_name = "PM-4405B"
    elif "HE-3301" in file: machine_name = "HE-3301"
    elif "BL-5702" in file: machine_name = "BL-5702"
    
    text_content = ""
    try:
        prs = Presentation(file)
        for slide in prs.slides:
            for shape in slide.shapes:
                if hasattr(shape, "text"):
                    text_content += shape.text + " "
    except Exception as e:
        print(f"Gagal membaca {file}: {e}")
        
    if text_content:
        knowledge_base.append({
            "machine": machine_name,
            "text": text_content,
            "file": os.path.basename(file)
        })

# Setup TF-IDF Vectorizer
tfidf_vectorizer = TfidfVectorizer(stop_words='english')
if knowledge_base:
    tfidf_matrix = tfidf_vectorizer.fit_transform([kb['text'] for kb in knowledge_base])

# ----------------------------------------------------
# 3. ENDPOINT API UNTUK DASHBOARD
# ----------------------------------------------------
@app.route('/api/predict', methods=['POST'])
def predict_anomaly():
    data = request.json
    machine = data.get('machine')
    sensor_data = data.get('sensor_data') # Input data realtime dari dashboard
    
    if machine not in models:
        return jsonify({"error": f"Model untuk {machine} belum di-training!"}), 400
        
    model = models[machine]
    features = cols[machine]
    
    # Konversi input ke DataFrame
    df_input = pd.DataFrame([sensor_data], columns=features)
    
    # --- PROSES AI 1: DETEKSI ANOMALI ---
    prediction = model.predict(df_input)[0]
    anomaly_score = model.decision_function(df_input)[0]
    
    is_anomaly = bool(prediction == -1)
    status = "CRITICAL (ANOMALY)" if is_anomaly else "NORMAL"
    
    response = {
        "machine": machine,
        "status": status,
        "anomaly_score": round(float(anomaly_score), 4),
        "explanation": [],
        "recommendation": {}
    }
    
    if is_anomaly:
        # --- PROSES AI 2: SHAP EXPLAINER ---
        explainer = shap.TreeExplainer(model)
        shap_values = explainer.shap_values(df_input)
        
        feature_contributions = []
        for i, col in enumerate(features):
            feature_contributions.append({
                "sensor": col,
                "value": float(df_input[col].iloc[0]),
                "contribution": float(shap_values[0][i])
            })
        
        # Urutkan penyebab (kontribusi paling besar)
        feature_contributions.sort(key=lambda x: abs(x['contribution']), reverse=True)
        response["explanation"] = feature_contributions
        
        # --- PROSES AI 3: TF-IDF RECOMMENDER ---
        top_sensor = feature_contributions[0]['sensor']
        query = f"{machine} anomaly {top_sensor} failure leak trip vibration temperature"
        
        if knowledge_base:
            query_vec = tfidf_vectorizer.transform([query])
            sim_scores = cosine_similarity(query_vec, tfidf_matrix).flatten()
            best_match_idx = sim_scores.argmax()
            
            if sim_scores[best_match_idx] > 0.05:
                best_match = knowledge_base[best_match_idx]
                response["recommendation"] = {
                    "source_rca": best_match["file"],
                    "match_confidence": round(float(sim_scores[best_match_idx]), 2),
                    "action": f"ACTION REQUIRED: Execute field inspection on {top_sensor} parameters in accordance with CAPA directives."
                }
            else:
                response["recommendation"] = {"message": "No matching historical anomaly patterns found in RCA database."}

    return jsonify(response)

if __name__ == '__main__':
    print("API Dashboard Berhasil Berjalan!")
    print("Frontend bisa POST data ke: http://localhost:5000/api/predict")
    app.run(debug=True, port=5000)
